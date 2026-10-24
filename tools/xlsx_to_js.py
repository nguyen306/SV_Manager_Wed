"""Chuyển data/students.xlsx thành data/students.js để trang web đọc được (kể cả khi mở bằng file://).

Cách chạy:  python tools/xlsx_to_js.py   (cần: pip install openpyxl)

File Excel gồm các sheet:
  1. Sheet đầu tiên: danh sách sinh viên + điểm hệ 4 từng môn (Năm 1..4) + GPA
  2. "TKB Chính Thức ...": thời khóa biểu ban ngày theo lớp
  3. "Lịch Phụ Đạo Tối ...": lịch phụ đạo buổi tối theo môn
  4. "Danh Sách SV Phụ Đạo Tối": sinh viên phải học phụ đạo (điểm môn < 2.0)
  5. "Cấu Hình Học Phí" + "Chi Tiết Học Phí": học phí theo năm (tạo bằng tools/build_tuition.py)
"""
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "students.xlsx"
OUT = ROOT / "data" / "students.js"

FIRST_GRADE_COL = 6  # cột G: môn đầu tiên của "Năm 1"
SUBJECT_RE = re.compile(r"\s*(.+?)\s*\((\d+)\)\s*$")  # "Giải tích(2)" -> ("Giải tích", 2)


def major_of(class_name):
    # "Công nghệ thông tin 2 K66" -> "Công nghệ thông tin"
    return re.sub(r"\s*\d*\s*K\d+$", "", class_name).strip()


def year_of(text):
    # "Năm 2" -> 2
    return int(re.search(r"\d+", str(text)).group())


def subject_name(text):
    m = SUBJECT_RE.match(str(text or ""))
    return m.group(1) if m else str(text).strip()


def find_sheet(wb, keyword):
    for ws in wb.worksheets:
        if keyword.lower() in ws.title.lower():
            return ws
    return None


def data_rows(ws, first_row=4):
    """Các dòng dữ liệu của sheet phụ (dòng 1 là tiêu đề lớn, dòng 3 là tên cột)."""
    if ws is None:
        return []
    return [r for r in ws.iter_rows(min_row=first_row, values_only=True) if any(c is not None for c in r)]


def read_students(ws):
    year_row = [c.value for c in ws[1]]     # dòng 1: "Năm 1", "Năm 2"... (ô gộp)
    subject_row = [c.value for c in ws[2]]  # dòng 2: "Tên môn(số tín chỉ)"
    gpa_col = year_row.index("GPA")

    subjects, year = [], None
    for col in range(FIRST_GRADE_COL, gpa_col):
        if year_row[col]:
            year = year_of(year_row[col])
        m = SUBJECT_RE.match(subject_row[col] or "")
        if m:
            subjects.append({"col": col, "name": m.group(1), "credits": int(m.group(2)), "year": year})

    students = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row[2]:
            continue
        diem = {s["name"]: row[s["col"]] for s in subjects if row[s["col"]] is not None}
        gpa = row[gpa_col]
        if not isinstance(gpa, (int, float)):
            # Cột GPA là công thức (trung bình có trọng số theo tín chỉ); nếu Excel chưa lưu
            # giá trị đã tính thì tự tính lại theo đúng công thức đó.
            graded = [s for s in subjects if s["name"] in diem]
            credits = sum(s["credits"] for s in graded)
            gpa = sum(diem[s["name"]] * s["credits"] for s in graded) / credits if credits else None
        students.append({
            "msv": str(int(row[2])),
            "hoTen": row[1].strip(),
            "lop": row[3].strip(),
            "nganh": major_of(row[3]),
            "email": row[4].strip(),
            "sdt": str(row[5]).strip(),
            "ngaySinh": None,  # file Excel chưa có cột ngày sinh
            "diem": diem,
            "gpa": round(gpa, 2) if gpa is not None else None,
        })
    return subjects, students


def read_tuition(wb, students):
    """Học phí theo năm. Các cột tính bằng công thức (đơn giá, phải nộp, thừa, nợ...) được tính lại
    ở đây theo đúng công thức trong Excel, vì openpyxl không đọc được kết quả công thức chưa lưu."""
    cfg_ws, detail_ws = find_sheet(wb, "Cấu Hình Học Phí"), find_sheet(wb, "Chi Tiết Học Phí")
    if cfg_ws is None or detail_ws is None:
        return None, []
    num = lambda v: float(v) if isinstance(v, (int, float)) else 0.0
    config = {"baseFee": num(cfg_ws["B3"].value), "yearlyIncrease": num(cfg_ws["B4"].value), "tutorFee": num(cfg_ws["B5"].value)}
    unit_price = lambda year: round(config["baseFee"] * (1 + config["yearlyIncrease"]) ** (year - 1))

    rows = []
    for r in data_rows(detail_ws):
        if not r[0]:
            continue
        year = year_of(r[3])
        subjects_count, tutor_count = int(num(r[5])), int(num(r[8]))
        exempt, paid, withdrawn = num(r[10]), num(r[12]), num(r[13])
        tuition = subjects_count * unit_price(year)
        tutor_fee = tutor_count * config["tutorFee"]
        payable = tuition + tutor_fee - exempt
        rows.append({
            "msv": str(r[0]).strip(), "year": year, "schoolYear": str(r[4] or ""),
            "subjects": subjects_count, "unitPrice": unit_price(year), "tuition": tuition,
            "tutorSubjects": tutor_count, "tutorFee": tutor_fee, "exempt": exempt, "payable": payable,
            "paid": paid, "withdrawn": withdrawn,
            "surplus": max(0, paid - withdrawn - payable), "debt": max(0, payable - (paid - withdrawn)),
            "dueDate": str(r[16] or ""), "receipt": str(r[17] or ""), "note": str(r[18] or ""),
        })

    # Đối chiếu tổng số môn / số môn phụ đạo của mỗi SV với sheet điểm
    totals = {}
    for x in rows:
        t = totals.setdefault(x["msv"], [0, 0])
        t[0] += x["subjects"]
        t[1] += x["tutorSubjects"]
    mismatch = [
        s["msv"] for s in students
        if totals.get(s["msv"]) != [len(s["diem"]), sum(1 for v in s["diem"].values() if v < 2.0)]
    ]
    if mismatch:
        print(f"Cảnh báo: {len(mismatch)} SV có số môn học phí không khớp sheet điểm, VD: {mismatch[:5]}")
    return config, rows


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    subjects, students = read_students(wb.worksheets[0])

    # Năm học, Lớp, Môn học, Thứ học, Ca học ban ngày
    timetable = [
        {"year": year_of(r[0]), "lop": r[1].strip(), "subject": subject_name(r[2]), "day": r[3].strip(), "session": r[4].strip()}
        for r in data_rows(find_sheet(wb, "TKB"))
    ]
    # Năm học, Môn học phụ đạo, Thứ học, Ca học buổi tối, Đối tượng tham gia
    tutoring_schedule = [
        {"year": year_of(r[0]), "subject": subject_name(r[1]), "day": r[2].strip(), "session": r[3].strip(), "target": (r[4] or "").strip()}
        for r in data_rows(find_sheet(wb, "Lịch Phụ Đạo"))
    ]
    # MSV, Họ và tên, Lớp, Năm học, Môn học cần phụ đạo, Điểm HP, Lịch học phụ đạo tối
    tutoring_students = [
        {"msv": str(int(r[0])), "year": year_of(r[3]), "subject": subject_name(r[4]), "score": r[5], "schedule": (r[6] or "").strip()}
        for r in data_rows(find_sheet(wb, "Danh Sách SV Phụ Đạo"))
    ]

    tuition_config, tuition = read_tuition(wb, students)

    # Kiểm tra tên môn ở các sheet phụ khớp với sheet điểm
    known = {s["name"] for s in subjects}
    unknown = {x["subject"] for x in timetable + tutoring_schedule + tutoring_students} - known
    if unknown:
        print("Cảnh báo: có môn không khớp với sheet điểm:", unknown)

    dump = lambda v: json.dumps(v, ensure_ascii=False, indent=1)
    OUT.write_text(
        "// Tự sinh từ data/students.xlsx bằng tools/xlsx_to_js.py — không sửa tay.\n"
        f"window.SUBJECTS = {dump([{k: s[k] for k in ('name', 'credits', 'year')} for s in subjects])};\n"
        f"window.STUDENTS = {dump(students)};\n"
        f"window.TIMETABLE = {dump(timetable)};\n"
        f"window.TUTORING_SCHEDULE = {dump(tutoring_schedule)};\n"
        f"window.TUTORING_STUDENTS = {dump(tutoring_students)};\n"
        f"window.TUITION_CONFIG = {dump(tuition_config)};\n"
        f"window.TUITION = {dump(tuition)};\n",
        encoding="utf-8",
    )
    print(
        f"Đã ghi {len(subjects)} môn học, {len(students)} sinh viên, {len(timetable)} dòng TKB, "
        f"{len(tutoring_schedule)} lịch phụ đạo, {len(tutoring_students)} lượt SV phụ đạo, "
        f"{len(tuition)} dòng học phí vào {OUT}"
    )


if __name__ == "__main__":
    main()
